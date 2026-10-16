"""ETL de la bandeja: Excel de Tesorería -> Postgres, y extracción de texto de PDFs."""
from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .extraccion import extraer_pdf


@dataclass(frozen=True)
class Doc:
    archivo: str
    sha256: str
    texto: str
    ruta: str = ""
    metodo: str = "nativo"
    confianza_ocr: float | None = None
    notas: tuple[str, ...] = ()


def leer_bandeja(carpeta: Path, conocidos: frozenset[str] | set[str] = frozenset()) -> list[Doc]:
    """Lee los PDF de la bandeja. Los ya procesados (por sha256) se omiten ANTES de extraer: no se repite OCR ni visión."""
    docs = []
    for p in sorted(carpeta.glob("*.pdf")):
        sha = hashlib.sha256(p.read_bytes()).hexdigest()
        if sha in conocidos:
            continue
        ext = extraer_pdf(p)
        notas = tuple(f"p{pg.numero} {pg.metodo}: {n}" for pg in ext.paginas for n in pg.notas)
        docs.append(Doc(p.name, sha, ext.texto, str(p), ext.metodo,
                        ext.confianza_min, notas))
    return docs


def _fecha(v) -> date:
    return v.date() if isinstance(v, datetime) else v if isinstance(v, date) else date.fromisoformat(str(v))


def cargar_registro_tesoreria(cx, xlsx: Path) -> int:
    ws = load_workbook(xlsx, data_only=True)["Garantias"]
    n = 0
    for numero, ordenante, benef, monto, moneda, vence, estado in list(ws.iter_rows(values_only=True))[1:]:
        if not numero:
            continue
        cx.execute(
            """INSERT INTO doc.registro_tesoreria(numero, ordenante, beneficiario, monto, moneda, fecha_vencimiento, estado_registro)
               VALUES (%s,%s,%s,%s,%s,%s,%s)
               ON CONFLICT (numero) DO UPDATE SET ordenante=EXCLUDED.ordenante, beneficiario=EXCLUDED.beneficiario,
                 monto=EXCLUDED.monto, moneda=EXCLUDED.moneda, fecha_vencimiento=EXCLUDED.fecha_vencimiento,
                 estado_registro=EXCLUDED.estado_registro""",
            (str(numero).strip(), ordenante, benef, Decimal(str(monto)), str(moneda).strip(), _fecha(vence), estado))
        n += 1
    return n
